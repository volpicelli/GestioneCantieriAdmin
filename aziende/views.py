from django.shortcuts import render
from django.http import HttpResponseRedirect,JsonResponse

from home.models import Azienda,Fornitori,Cliente,BancaFornitori,CondizioniPagamento
from home.form_aziende import FormAzienda
from django.views.generic.list import ListView,View
from django.views.generic.edit import CreateView,UpdateView
from django.urls import reverse_lazy
from django.shortcuts import redirect
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
import json,os
# Create your views here.

class AziendeList(ListView):
    model = Azienda
    template_name = "listAziende.html"

    def get(self,request):

        if not request.user.is_authenticated:        
            return HttpResponseRedirect('/access/login')

        az = Azienda.objects.all()

        context={'aziende': az}
        

        return render(request, self.template_name, context)
       
class AziendaInsert(CreateView):
    model = Azienda
    template_name = "insertazienda.html"

    def get(self,request):
        if not request.user.is_authenticated:        
            return HttpResponseRedirect('/access/login')
        form = FormAzienda()
    
        return render(request, self.template_name, {'form': form})
    

    def post(self,request):

        res = request.POST
        res2={}
        res2['post']= res
        filelogo = False
        if 'logo' in request.FILES:
            file = request.FILES.get('logo')
            res2['file']=file.name
            filelogo=True
        else:
            res2['file']='No File logo'

        codcf = request.POST.get('codcf')
        nome = request.POST.get('nome')
        ragione_sociale=request.POST.get('ragione_sociale',None)
        codfisc=request.POST.get('codfisc',None)
        descrizione =request.POST.get('descrizione',None)
        indirizzo=request.POST.get('indirizzo',None)
        cap=request.POST.get('cap',None)
        local=request.POST.get('local',None)
        prov=request.POST.get('prov',None)
        partiva=request.POST.get('partiva',None)
        telefono=request.POST.get('telefono',None)
        cellulare=request.POST.get('cellulare',None)
        pec=request.POST.get('pec',None)
        fax=request.POST.get('fax',None)
        email=request.POST.get('email',None)
        resprap=request.POST.get('resprap',None)
        fmemo=request.POST.get('fmemo',None)
        nome_pf=request.POST.get('nome_pf',None)
        cogn_pf=request.POST.get('cogn_pf',None)

        a= Azienda(codcf=codcf,nome=nome)
        a.save()
        if filelogo:
            a.logo=file
            a.save()
        az = Azienda.objects.filter(codcf=codcf)
        # Devo inserire l'azienda anche tra i fornitori.
        f = Fornitori(codcf=codcf,azienda=a,fmemo=" Azienda Fornitore di se stessa")
        f.save()
        al =list(az.values())[0]
        #s = serializers.serialize("json",az)
        return HttpResponseRedirect(request.META.get('HTTP_REFERER', '/'))      

        #return redirect('.')


       
class UpdateAzienda(UpdateView):
    model = Azienda
    template_name = "insertazienda.html"
    form_class = FormAzienda
    success_url = reverse_lazy('listclienti')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['page_title'] = 'Update Book'
        return context
    
    # Optional: You can customize the form before it's rendered
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        # Customize form here if needed
        return form
    
    # Optional: You can perform additional actions before saving
    def form_valid(self, form):
        # You can access the object being updated with self.object
        # or form.instance before it's saved
        return super().form_valid(form)

    #def get(self,request,pk):
    #    form = FormClientiGestione()
    
    #    return render(request, 'nuovocliente.html', {'form': form})
 
#met = MyExifTool(f.temporary_file_path())

class LoadFornitori(View):
    def get(self,request):
        if not request.user.is_authenticated:        
            return HttpResponseRedirect('/access/login')
        #form = FormAzienda()
        az = Azienda.objects.all()
        context={'aziende': az}
    
        return render(request, 'loadfornitori.html', context)
    
    def post(self,request):
        res= request.POST
        file = request.FILES.get('file')

        res2={}
        res2['file']=file.name
        #res2['post']=res
        file_uploaded = request.FILES.get('file')
        tabella = request.POST.get('tabella')
        extension = os.path.splitext(file_uploaded.name)[1]
        if 'json' not in extension:
            res2['messaggio']= "File non valido, caricare un file json"
            return JsonResponse(res2,safe=False)
        if file_uploaded:
            data = json.loads(file_uploaded.read())
            if tabella == 'condizionipagamento':
                azienda = request.POST.get('azienda')

                try:
                    cp=CondizioniPagamento.objects.filter(azienda_id=azienda)
                except:
                    cp=None
                if len(cp)>1:
                    res2['messaggio']= "Condizioni di pagamento sono gia` presenti per questa azienda"
                    return JsonResponse(res2,safe=False)

                for one in data:
                    one['azienda_id']=azienda
                    c = CondizioniPagamento(**one)
                    c.save()

                res2['messaggio']= "Condizioni di pagamento caricati con successo"
                return JsonResponse(res2,safe=False)   
            if tabella == 'bancafornitori':
                azienda = request.POST.get('azienda')
                """
                try:
                    bfa=BancaFornitori.objects.all()
                except:
                    bfa=None
                if len(bfa)>1:
                    res2['messaggio']= "Banca fornitori sono gia` presenti"
                    return JsonResponse(res2,safe=False)
                """
                for one in data:

                    try:
                        forn = Fornitori.objects.filter(azienda_id=azienda)
                    except:
                        forn = None
                        res2['messaggio']= "Nessun fornitore trovato per l'azienda selezionata"
                        return JsonResponse(res2,safe=False)
                    if len(forn) ==0 :
                        res2['messaggio']= "Nessun fornitore trovato per l'azienda selezionata"
                        return JsonResponse(res2,safe=False)

                    try:
                        f = forn.get(codcf=one['codfor'])  # Fornitore 
                    except:
                        f = None
            

                    bf = BancaFornitori(**one)
                    if f:
                        bf.fornitore_id=f.id
                    else:
                        bf.fornitore_id=f
                    #print(one)
                    bf.save()
                res2['messaggio']= "Banche Fornitori  caricati con successo"
                return JsonResponse(res2,safe=False)   
            if tabella == 'fornitori':
                azienda = request.POST.get('azienda')

                for one in data:
                    one.pop('azienda')
                    try:
                        codpag = CondizioniPagamento.objects.get(codpag=one['codpag'])
                    except: 
                        codpag = None
                    one.pop('codpag')
                    one['azienda_id']=azienda
                    f = Fornitori(**one)
                    f.codpag=codpag
                    f.save()
                res2['messaggio']= "Caricamento avvenuto con successo di "+str(len(data))+" records"
                return JsonResponse(res2,safe=False)
            if tabella == 'clienti':
                azienda = request.POST.get('azienda')

                for one in data:
                    one.pop('azienda')
                    try:
                        codpag = CondizioniPagamento.objects.get(codpag=one['codpag'])
                    except: 
                        codpag = None
                    #one.pop('codpag')
                    one['azienda_id']=azienda
                    f = Cliente(**one)
                    f.codpag=codpag
                    f.save()
                res2['messaggio']= "Caricamento avvenuto con successo di "+str(len(data))+" records"
                return JsonResponse(res2,safe=False)
                    
            #try:
            #    codpag = CondizioniPagamento.objects.get(codpag=one['codpag'])
            #except: 
            #    codpag = None
            #one.pop('codpag')
#
            #f = Cliente(**one)
            #f.codpag=codpag
            #one['codpag'] = codpag
            #print(one)
            #f.save()


            """
            work_book = load_workbook(file_uploaded)
            events_sheet = work_book.active 
            # Get fields name array 
            fieldnames=[]
            res2['fieldnames']=[]
            for line in range(1,2): # events_sheet.max_row+1):
                for col in events_sheet.iter_cols(1, events_sheet.max_column):
                    fieldnames.append(col[line].value)
                    
            items=[]
            for line in range(2, events_sheet.max_row):
                idxcol=0
                row={'azienda_id':azienda}
                for col in events_sheet.iter_cols(1, events_sheet.max_column):
                    row[fieldnames[idxcol]]=col[line].value
                    idxcol+=1

                print(row)
                items.append(row)
            """

            """
            if tabella == 'fornitori':
                for one in items:
                    f = Fornitori(**one)
                    f.save()
            else:
                for one in items:
                    c = Cliente(**one)
                    c.save()
            """
        

